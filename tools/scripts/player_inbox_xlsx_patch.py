#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""玩家持久收件箱(docs/design/player-inbox-service.md §5.3)第 1 批要对 data/tip/Tip.xlsx 做的改动,写成可重复执行的脚本。

为什么是脚本而不是手改(同 tools/scripts/equip_xlsx_patch.py):
  - Tip.xlsx 同时有多个会话在写,它是二进制、不能 3-way 合并,整表写回会**静默**删掉别人的行
    (导表器照常跑通,只是少发几个号)。唯一安全的写法是 openpyxl 从磁盘重新载入 → 只往**空格子**里写
    自己那几行 → save → 读回来逐格核对,再按导表器的口径(组 / 码 / fault)逐组核对一遍。
  - 本脚本对既有表的每一次写入都先断言目标格为空,所以「别人的格子一格不动」是构造保证,不是约定。
  - 新组追加在表尾(组头前空一行,与既有组一致)。表尾之后不存在「别人的组」,所以不需要插入行;
    若重跑时本组之后已有别人追加的组,幂等判定只看本组开头的码,照样报「已是目标状态」。

子命令(幂等:已是目标状态就只报告、不改文件;发现与预期不符的现状就中止,绝不覆盖):

    py -3 tools/scripts/player_inbox_xlsx_patch.py tip-codes             # 新开 //player_inbox_error 段 + 15 个码
    py -3 tools/scripts/player_inbox_xlsx_patch.py tip-codes --dry-run   # 只打印将要做什么,不写文件

写之前的只读核对(任一不满足即中止,什么都不写):
  - 第 1 行表头是 name / commen / … / fault(fault 恰好一列);
  - 15 个码名在表里全局唯一,也不在发号 state 的任何组 / legacy_ids 里(生成的 tip proto 没有 package);
  - [29000, 30000) 段没被表里任何组占用、没被说明行预留、没被发号 state(含墓碑组)占用,
    也不在已生成的 go/shared/generated/tip/segments.go 里。

执行前先确认 `git status --short data/tip/Tip.xlsx` 没有别人的未提交改动(任一时刻只允许一方有未提交修改)。
执行后由 Codex 导表(`dev.bat export`),核对 segments.go 出现 player_inbox_error / Base 29000 的段、
faults.go 的新段只有 PlayerInboxStoreUnavailable 一个码,再把新生成的 proto/tip/player_inbox_error_tip.pb.{h,cc}
登记进 cpp/generated/table 的 CMakeLists.txt / table.vcxproj / table.vcxproj.filters(AGENTS §4.4)。
需要 Python 3 + openpyxl(与导表器同一套依赖)。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - 环境问题,给一句能照着做的话
    sys.exit("缺少 openpyxl:请先执行  pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parents[2]
TIP_XLSX = REPO_ROOT / "data" / "tip" / "Tip.xlsx"
TIP_STATE = REPO_ROOT / "tools" / "data_table_exporter" / "state" / "mapping" / "tip_enum_ids" / "tip_enum_ids.json"
TIP_SEGMENTS_GO = REPO_ROOT / "go" / "shared" / "generated" / "tip" / "segments.go"

HEADER_ROW = 1
TIP_DATA_BEGIN_ROW = 18            # 与导表器 enum_gen._TIP_DATA_BEGIN_ROW 一致
TIP_PREFIX = "//"


# ════════════════════════════════════════════════════════════════════════════
# 数据常量表(设计 §5.3)。改码只改这一段。
# ════════════════════════════════════════════════════════════════════════════

# A 列码名不带 k(导表器生成 PlayerInboxError_kXxx);全部带 PlayerInbox 前缀,全局唯一。
# 行序就是号序:导表器按表内出现顺序在本段取最小空位,所以这里的顺序上线后不要调整,只往末尾追加。
# fault=True 的是服务端内部故障(库或核心 Redis 不可用),其余都是业务拒绝 / 协议拒绝,fault 列留空。
# 文案标点用半角,与 equip_error / pet_error 段一致。
TIP_GROUP = "player_inbox_error"
TIP_BASE = 29000
TIP_WIDTH = 1000
TIP_CODES = [
    # 码名                               文案                                   fault
    # —— 入箱前置拒绝 ——
    ("PlayerInboxDisabled",              "该类收件箱消息暂未开放",               False),
    ("PlayerInboxUnknownType",           "未知的收件箱消息类型",                 False),
    ("PlayerInboxPlayerNotFound",        "玩家不存在",                           False),
    # —— 入箱校验 ——
    ("PlayerInboxPayloadInvalid",        "收件箱消息内容无效",                   False),
    ("PlayerInboxPayloadTooLarge",       "收件箱消息内容过大",                   False),
    ("PlayerInboxTextTooLong",           "操作人、理由或备注过长",               False),
    ("PlayerInboxIssuedOutOfWindow",     "签发时间超出允许范围",                 False),
    ("PlayerInboxFull",                  "该玩家的收件箱已满",                   False),
    ("PlayerInboxIdempotencyConflict",   "请求号已被不同内容使用",               False),
    # —— 围栏(后两个同时告警,但属于拒绝而非服务端故障)——
    ("PlayerInboxStaleOwner",            "场景归属已过期",                       False),
    ("PlayerInboxEpochRequired",         "缺少场景归属代际",                     False),
    ("PlayerInboxEpochAhead",            "场景归属代际超前",                     False),
    # —— 运维接口 ——
    ("PlayerInboxItemNotFound",          "收件箱条目不存在",                     False),
    ("PlayerInboxItemStateConflict",     "收件箱条目当前状态不允许该操作",       False),
    # —— 服务端故障 ——
    ("PlayerInboxStoreUnavailable",      "收件箱服务暂不可用,请稍后再试",        True),
]


# ════════════════════════════════════════════════════════════════════════════
# 通用工具(口径与 equip_xlsx_patch.py 相同;本脚本刻意不 import 它,两份一次性补丁脚本互不耦合)
# ════════════════════════════════════════════════════════════════════════════

Patch = tuple[int, int, object]     # (行, 列, 值),行列都从 1 起


def fail(message: str) -> "NoReturn":
    sys.exit(f"[中止] {message}")


def rel(path: Path) -> str:
    return path.relative_to(REPO_ROOT).as_posix()


def only_sheet(wb, path: Path):
    if len(wb.sheetnames) != 1:
        fail(f"{path.name} 预期只有一个 sheet,实际 {wb.sheetnames}")
    return wb[wb.sheetnames[0]]


def sheet_matrix(ws) -> list[list]:
    return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]


def _trimmed(matrix: list[list]) -> dict[tuple[int, int], object]:
    """矩阵 -> {(行, 列): 值},只留非空格。两张表「非空格完全相同」即内容相同,与尾部留白多少无关。"""
    return {(r, c): v for r, row in enumerate(matrix, start=1) for c, v in enumerate(row, start=1)
            if v is not None}


def first_difference(actual: list[list], expected: list[list]) -> str | None:
    a, e = _trimmed(actual), _trimmed(expected)
    if a == e:
        return None
    r, c = min(k for k in set(a) | set(e) if a.get(k) != e.get(k))
    return f"{get_column_letter(c)}{r}:读回 {a.get((r, c))!r},预期 {e.get((r, c))!r}"


def apply_patches_and_verify(xlsx_path: Path, wb, ws, before: list[list], patches: list[Patch]) -> None:
    """只往空格子里写 → save → 读回逐格核对。「目标格必须为空」保证别人的格子一格都不会被改,
    也保证重复执行不会把同一批数据叠写两遍。"""
    for row, col, value in patches:
        if value is None:
            fail(f"内部错误:补丁里出现了空值({get_column_letter(col)}{row})")
        if ws.cell(row=row, column=col).value is not None:
            fail(f"{xlsx_path.name} 的 {get_column_letter(col)}{row} 已有内容 "
                 f"{ws.cell(row=row, column=col).value!r},本脚本只写空格、不覆盖 —— 先查清是谁写的。")
    for row, col, value in patches:
        ws.cell(row=row, column=col, value=value)
    wb.save(xlsx_path)

    width = max([len(r) for r in before] + [c for _r, c, _v in patches])
    height = max([len(before)] + [r for r, _c, _v in patches])
    expected = [list(r) + [None] * (width - len(r)) for r in before]
    expected += [[None] * width for _ in range(height - len(expected))]
    for row, col, value in patches:
        expected[row - 1][col - 1] = value

    check_ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path)
    diff = first_difference(sheet_matrix(check_ws), expected)
    if diff:
        fail(f"保存后读回的 {xlsx_path.name} 与预期不符({diff})—— 立刻 git checkout -- {rel(xlsx_path)} 还原")


# ── Tip.xlsx 解析(判据与导表器 enum_gen._parse_tip_sheet 一致)──────────────

def is_group_header(name: str) -> bool:
    """"//" 后**紧贴**非空白字符 = 组头,"// 这是一句说明" = 说明行。不能用"有没有 base="来分。"""
    return name.startswith(TIP_PREFIX) and len(name) > 2 and not name[2].isspace()


def parse_group_header(name: str, row_idx: int) -> tuple[str, int, int]:
    parts = name[2:].split()
    kv = dict(tok.split("=", 1) for tok in parts[1:] if "=" in tok)
    try:
        return parts[0], int(kv["base"]), int(kv["width"])
    except (KeyError, ValueError):
        fail(f"{TIP_XLSX.name} 第 {row_idx} 行 {name!r} 像组头却没有合法的 base= / width= —— 导表器也会拒绝,先修表")


def tip_parse(matrix: list[list], fault_col: int):
    """返回 (组列表, 说明行列表)。组 = dict(name, base, width, row, codes=[(码名, 文案, fault 原值)],
    code_rows=[与 codes 逐项对应的行号])。"""
    groups: list[dict] = []
    notes: list[tuple[int, str]] = []
    for row_idx in range(TIP_DATA_BEGIN_ROW, len(matrix) + 1):
        row = matrix[row_idx - 1]
        a = row[0]
        if not isinstance(a, str) or not a.strip():
            continue
        name = a.strip()
        if name.startswith(TIP_PREFIX):
            if not is_group_header(name):
                notes.append((row_idx, name))
                continue
            group, base, width = parse_group_header(name, row_idx)
            groups.append(dict(name=group, base=base, width=width, row=row_idx, codes=[], code_rows=[]))
            continue
        if not groups:
            fail(f"{TIP_XLSX.name} 第 {row_idx} 行的码 {name!r} 出现在第一个组头之前")
        fault = row[fault_col - 1] if fault_col <= len(row) else None
        groups[-1]["codes"].append((name, row[1] if len(row) > 1 else None, fault))
        groups[-1]["code_rows"].append(row_idx)
    return groups, notes


def generated_segments() -> list[tuple[str, int, int]]:
    """只读解析已生成的 segments.go:[(组名, base, width)]。文件不存在就返回空(新检出尚未导表)。"""
    if not TIP_SEGMENTS_GO.exists():
        return []
    pattern = re.compile(r'Group:\s*"([^"]+)",\s*Base:\s*(\d+),\s*Width:\s*(\d+)')
    return [(m.group(1), int(m.group(2)), int(m.group(3)))
            for m in pattern.finditer(TIP_SEGMENTS_GO.read_text(encoding="utf-8"))]


def self_check() -> None:
    names = [name for name, _t, _f in TIP_CODES]
    if len(set(names)) != len(names) or len(names) > TIP_WIDTH:
        fail("TIP_CODES 码名重复或超过段宽")
    if any(not name.startswith("PlayerInbox") or not name.isidentifier() or name.startswith("k") for name in names):
        fail("TIP_CODES 码名必须是合法标识符、以 PlayerInbox 开头、不带 k 前缀")
    if [name for name, _t, fault in TIP_CODES if fault] != ["PlayerInboxStoreUnavailable"]:
        fail("TIP_CODES 里只有 PlayerInboxStoreUnavailable 是服务端故障(设计 §5.3)")
    if any(not text or not text.strip() for _n, text, _f in TIP_CODES):
        fail("TIP_CODES 文案不能为空")


# ════════════════════════════════════════════════════════════════════════════
# tip-codes
# ════════════════════════════════════════════════════════════════════════════

def cmd_tip_codes(dry_run: bool) -> None:
    wb = openpyxl.load_workbook(TIP_XLSX)
    ws = only_sheet(wb, TIP_XLSX)
    before = sheet_matrix(ws)
    header = [str(v).strip().lower() if v is not None else "" for v in before[HEADER_ROW - 1]]
    if header[:2] != ["name", "commen"] or header.count("fault") != 1:
        fail(f"{TIP_XLSX.name} 第 1 行应为 name / commen / … / fault(fault 恰好一列),实际 {before[HEADER_ROW - 1]}")
    fault_col = header.index("fault") + 1

    groups, notes = tip_parse(before, fault_col)
    header_text = f"//{TIP_GROUP} base={TIP_BASE} width={TIP_WIDTH}"
    wanted = [(name, text, 1 if fault else None) for name, text, fault in TIP_CODES]
    new_names = [name for name, _t, _f in TIP_CODES]

    mine = next((g for g in groups if g["name"] == TIP_GROUP), None)
    if mine is not None:
        # 已是目标状态 = 段声明一致,且这 15 个码是本组开头的连续一段。不要求恰好在组尾:以后还会往组尾追加。
        if (mine["base"], mine["width"]) == (TIP_BASE, TIP_WIDTH) and mine["codes"][:len(wanted)] == wanted:
            print(f"[已是目标状态] {header_text}(第 {mine['row']} 行)下已有这 {len(wanted)} 个码"
                  f"(码名 / 文案 / fault 一致),不改文件。")
            return
        fail(f"{TIP_XLSX.name} 第 {mine['row']} 行已有 //{TIP_GROUP} 组,但段声明或开头 {len(wanted)} 个码与本脚本不同"
             f"(可能只写进了一部分,或被手改过)—— 先查清是谁改的,本脚本不猜。")

    # 以下都是「先核对再动手」:码名全表唯一、29000 段没被表 / 说明行 / 发号 state / 已生成段表占用。
    owner = {code: g["name"] for g in groups for code, _t, _f in g["codes"]}
    clash = [f"{n}(在 {owner[n]})" for n in new_names if n in owner]
    if clash:
        fail(f"这些码名已在表里:{clash} —— 码名全局唯一(生成的 tip proto 没有 package),先查清。")
    lo, hi = TIP_BASE, TIP_BASE + TIP_WIDTH
    overlap = [f"{g['name']} [{g['base']},{g['base'] + g['width']})" for g in groups
               if g["base"] < hi and lo < g["base"] + g["width"]]
    if overlap:
        fail(f"段 [{lo},{hi}) 已被表里的组占用:{overlap} —— 改用下一个空闲整千段并同步设计 §5.3。")
    reserved = [f"第 {row} 行 {text!r}" for row, text in notes if str(TIP_BASE) in text]
    if reserved:
        fail(f"{TIP_BASE} 段已被说明行预留:{reserved} —— 改用下一个空闲整千段并同步设计 §5.3。")
    if TIP_STATE.exists():
        # 发号 state 里的整组是墓碑:表里删掉的组仍占着段。只读核对,不写 state。
        state = json.loads(TIP_STATE.read_text(encoding="utf-8"))
        state_groups = state.get("groups", {})
        tomb = [f"{n} [{g['base']},{g['base'] + g['width']})" for n, g in state_groups.items()
                if int(g["base"]) < hi and lo < int(g["base"]) + int(g["width"])]
        if tomb:
            fail(f"发号 state({rel(TIP_STATE)})里段 [{lo},{hi}) 已属于 {tomb} —— 换一个段,不要手改 state。")
        state_names = {name: g_name for g_name, g in state_groups.items() for name in g.get("ids", {})}
        state_names.update({entry.get("name"): f"legacy:{entry.get('group')}"
                            for entry in state.get("legacy_ids", {}).values() if isinstance(entry, dict)})
        state_clash = [f"{n}(在 {state_names[n]})" for n in new_names if n in state_names]
        if state_clash:
            fail(f"这些码名已在发号 state 里(可能是墓碑):{state_clash} —— 码名全局唯一,换名字,不要手改 state。")
    seg_overlap = [f"{name} [{base},{base + width})" for name, base, width in generated_segments()
                   if base < hi and lo < base + width]
    if seg_overlap:
        fail(f"已生成的 {rel(TIP_SEGMENTS_GO)} 里段 [{lo},{hi}) 已属于 {seg_overlap} —— 换一个段。")

    last_row = max((r for r, row in enumerate(before, start=1) if any(v is not None for v in row)), default=0)
    header_row = last_row + 2          # 与既有组一样,组头前留一个空行
    patches: list[Patch] = [(header_row, 1, header_text)]
    print(f"{TIP_XLSX.name}:现有 {len(groups)} 个组、{len(owner)} 个码;最后一个非空行是第 {last_row} 行;"
          f"段 [{lo},{hi}) 空闲;{len(new_names)} 个码名全局唯一。将追加:")
    print(f"  第 {header_row} 行: {header_text}")
    for offset, (name, text, fault) in enumerate(wanted, start=1):
        row_idx = header_row + offset
        patches += [(row_idx, 1, name), (row_idx, 2, text)]
        if fault is not None:
            patches.append((row_idx, fault_col, fault))
        print(f"  第 {row_idx} 行: {name} | {text}" + (" | fault=1" if fault else ""))
    if dry_run:
        print("[dry-run] 未写文件。")
        return
    apply_patches_and_verify(TIP_XLSX, wb, ws, before, patches)

    # 再按导表器的口径(组 / 码 / fault)核一遍:fault 标记若错了行,落到相邻的码上就是静默改了故障分类。
    check_ws = only_sheet(openpyxl.load_workbook(TIP_XLSX), TIP_XLSX)
    groups_after, notes_after = tip_parse(sheet_matrix(check_ws), fault_col)
    expected_groups = [(g["name"], g["base"], g["width"], g["codes"]) for g in groups] \
        + [(TIP_GROUP, TIP_BASE, TIP_WIDTH, wanted)]
    if [(g["name"], g["base"], g["width"], g["codes"]) for g in groups_after] != expected_groups \
            or [text for _row, text in notes_after] != [text for _row, text in notes]:
        fail(f"保存后各组的码与预期不符 —— 立刻 git checkout -- {rel(TIP_XLSX)} 还原")
    print(f"[完成] 组 {len(groups)} -> {len(groups_after)};既有 {len(owner)} 个码(码名, 文案, fault)逐项不变。"
          f"导表后核对:这 {len(wanted)} 个码的 id 应为 {lo}..{lo + len(wanted) - 1};tip_enum_ids.json 只增不改;"
          f"faults.go 新段只有 PlayerInboxStoreUnavailable;把新生成的 proto/tip/{TIP_GROUP}_tip.pb.{{h,cc}} "
          f"登记进 cpp/generated/table 的 CMakeLists.txt / table.vcxproj / table.vcxproj.filters。")


COMMANDS = {
    "tip-codes": cmd_tip_codes,
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=tuple(COMMANDS))
    parser.add_argument("--dry-run", action="store_true", help="只打印将要做什么,不写文件")
    args = parser.parse_args()
    self_check()
    COMMANDS[args.command](args.dry_run)


if __name__ == "__main__":
    # Windows 默认控制台编码(cp936 / cp1252)打不出部分字符时会在 print 上抛 UnicodeEncodeError,
    # 而那时文件可能已经写完 —— 统一成 utf-8,免得"改成功了却以非 0 退出"误导人。
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    main()
