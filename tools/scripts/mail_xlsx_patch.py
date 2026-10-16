#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""邮件系统 M1a(docs/design/mail-system.md §5.5 / §5.6)要对两张 xlsx 做的改动,写成可重复执行的脚本。

为什么是脚本而不是手改:
  - data/tip/Tip.xlsx 与 data/MessageLimiter.xlsx 同时有多个会话在写,它们是二进制、不能 3-way 合并,
    整表写回会**静默**删掉别人的行(导表器照常跑通,只是少发几个号 / 少几行档位)。
    唯一安全的写法是 openpyxl 从磁盘重新载入 → 只动自己那几格 → save → 读回来逐格核对。
  - MessageLimiter 按**消息号**登记档位,而新方法落到哪个号要等全量 proto-gen 之后才定,
    所以 message-limiter 按方法名去 proto/message_id.txt 查号,不把号写死。

两个子命令都**幂等**:已经是目标状态就只报告、不改文件;发现与预期不符的现状就中止,绝不覆盖。

用法(仓库根目录执行;需要 Python 3 + openpyxl + pyyaml,与导表器同一套依赖):

    python tools/scripts/mail_xlsx_patch.py tip-segment        # ① Tip.xlsx 新开 //mail_error 段 + 8 个码
    #   —— 第一次导表 → 全量 proto-gen(先 proto-gen-build,归构建方;之后跑 tools/scripts/mail_postgen.py verify)——
    python tools/scripts/mail_xlsx_patch.py message-limiter    # ② 全量 proto-gen 之后、第二次导表之前

完整的生成后步骤(含 C++ 工程登记与 robot vendor 同步)与通过标准见 docs/design/mail-system.md §12。

加 --dry-run 只打印将要做什么,不写文件。执行前先确认
`git status --short data/tip/Tip.xlsx data/MessageLimiter.xlsx` 没有别人的未提交改动。

message-limiter 按**数字消息号**往二进制表里写行:只在「最终要进 main 的那棵树」上、消息号已经定死之后跑。
号一旦重发,先前写下的行会留在旧号上,变成给别的 RPC 限流(脚本只认得出「那个号上已有一行不同的档位」,
认不出「这一行原本是我写的」)。

本脚本对既有表格的写入只有两种,都先核对现状:
  - 往**空格子**里追加新行(组头 + 8 个码;5 行档位)——目标格非空即中止;
  - 把 Tip.xlsx 里那一行**预留说明**改写一次(从清单里划掉已建组的 mail_error)——该格必须逐字等于已知的旧文案。
除此之外不改任何既有单元格。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - 环境问题,给一句能照着做的话
    sys.exit("缺少 openpyxl:请先执行  pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = REPO_ROOT / "data"
SCHEMA_DIR = DATA_DIR / "schema"
EXPORTER_DIR = REPO_ROOT / "tools" / "data_table_exporter"

# 主表版式(data/AGENTS.md):第 1 行列名 / 第 2-5 行表头区 / 第 6 行起数据。
HEADER_ROW = 1
DATA_BEGIN_ROW = 6


# ════════════════════════════════════════════════════════════════════════════
# 数据常量(设计文档 §5.5 / §5.6)。改内容只改这一段。
# ════════════════════════════════════════════════════════════════════════════

# §5.6 tip 码(A 列不带 k;导表器生成 MailError_kMailXxx)。mail 段**没有 fault 码**:
# 发件方被拒是规则生效,不是我方出错;存储 / 号段故障复用 common 段的 kServiceUnavailable。
TIP_MAIL_GROUP = "mail_error"
TIP_MAIL_BASE = 17000
TIP_MAIL_WIDTH = 1000
TIP_MAIL_CODES = [
    # 码名                      文案
    ("MailClaimNotOpen",        "附件领取暂未开放"),
    ("MailNotFound",            "邮件不存在或已删除"),
    ("MailExpired",             "邮件已过期"),
    ("MailRecipientBoxFull",    "收件人邮箱已满"),
    ("MailAttachmentNotOpen",   "当前环境不允许发送带附件的邮件"),
    ("MailAttachmentInvalid",   "邮件附件不合法"),
    ("MailRequestConflict",     "发件请求编号重复但内容不同"),
    ("MailSystemMailLimit",     "有效系统邮件数量已达上限"),
]

# Tip.xlsx 里为移植域预留号段的那一行说明(2026-10-13 位于第 258 行,行号只作参考,脚本按内容找)。
# mail_error 建组后把它从清单里划掉,免得后来的人以为 17000 还空着;其余两个域的预留原样保留。
TIP_RESERVATION_OLD = "//   mail_error 17000 / chat_error 18000 / auction_error 19000"
TIP_RESERVATION_NEW = "//   chat_error 18000 / auction_error 19000(mail_error 17000 已建组,在表尾)"

# §5.5 限频档位(次/秒)。NotifyMailEvent 是 S2C 下行,客户端不发,按惯例不进表;MailAdmin 不经 gate,不受此表约束。
LIMITER_SERVICE = "ClientPlayerMail"
LIMITER_METHODS = {
    "ListMails": 5,                 # 写后再拉 + 翻页
    "ReadMail": 10,                 # 客户端"选中即读",连续点选
    "DeleteMails": 3,               # 写操作(= gate 默认档,仍显式配)
    "DeleteReadMails": 3,
    "ClaimMailAttachments": 3,
}
LIMITER_S2C_METHOD = "NotifyMailEvent"     # 只用来确认 message_id.txt 确实出自含 mail 的那次 proto-gen
LIMITER_TIME_WINDOW = 1
# tip_message 列照抄表内既有行(全部是 1000)。这一列目前**没有读者**:限流器
# (cpp/libs/engine/core/message_limiter/message_limiter.cpp)硬编码回 kRateLimitExceeded,不读表。
# 1000 实际是 kSuccess(kRateLimitExceeded 是 1008),表内既有各行都是这个值;在有人给这一列接上读者、
# 并把全表一次改对之前,新行与既有行保持一致,不单独换口径(登记在 docs/design/mail-system.md §12)。
LIMITER_TIP_MESSAGE = 1000


# ════════════════════════════════════════════════════════════════════════════
# 通用工具
# ════════════════════════════════════════════════════════════════════════════

def fail(message: str) -> "NoReturn":
    sys.exit(f"[中止] {message}")


def rel(path: Path) -> str:
    return path.relative_to(REPO_ROOT).as_posix()


def file_digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def only_sheet(wb, path: Path, expect_title: str | None = None):
    if len(wb.sheetnames) != 1:
        fail(f"{path.name} 预期只有一个 sheet,实际 {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    if expect_title is not None and ws.title != expect_title:
        fail(f"{path.name} 的 sheet 名应为 {expect_title!r}(导表器按它匹配 schema),实际 {ws.title!r}")
    return ws


def sheet_matrix(ws) -> list[list]:
    return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]


def _trimmed(matrix: list[list]) -> dict[tuple[int, int], object]:
    """矩阵 -> {(行, 列): 值},只留非空格。两张表「非空格完全相同」即内容相同,与尾部留白多少无关。"""
    return {(r, c): v for r, row in enumerate(matrix, start=1) for c, v in enumerate(row, start=1)
            if v is not None}


def first_difference(actual: list[list], expected: list[list]) -> str | None:
    a, e = _trimmed(actual), _trimmed(expected)
    if a == e:
        return None
    r, c = min(k for k in set(a) | set(e) if a.get(k) != e.get(k))
    return f"{get_column_letter(c)}{r}:读回 {a.get((r, c))!r},预期 {e.get((r, c))!r}"


Patch = tuple[int, int, object]     # (行, 列, 值),行列都从 1 起
Rewrite = tuple[int, int, object, object]   # (行, 列, 必须逐字相等的旧值, 新值)


def apply_and_verify(xlsx_path: Path, sheet: str | None, wb, ws, before: list[list],
                     patches: list[Patch], rewrites: list[Rewrite] = (), loaded_digest: str = "") -> None:
    """patches 只往空格子里写;rewrites 只改「现值逐字等于旧值」的格子 → save → 读回逐格核对。

    这两条前置检查是本脚本对既有表的全部写入纪律:别人的格子一格都不会被改,
    重复执行也不会把同一批数据叠写两遍。loaded_digest 是载入时文件的 SHA-256:保存前再算一次,
    不同 = 载入之后别的会话改过这张表,保存会把它的改动整表抹掉 —— 中止,重跑即可。
    """
    for row, col, value in patches:
        if value is None:
            fail(f"内部错误:补丁里出现了空值({get_column_letter(col)}{row})")
        current = ws.cell(row=row, column=col).value
        if current is not None:
            fail(f"{xlsx_path.name} 的 {get_column_letter(col)}{row} 已有内容 {current!r},"
                 f"本脚本只写空格、不覆盖 —— 先查清是谁写的。")
    for row, col, old, new in rewrites:
        current = ws.cell(row=row, column=col).value
        if current != old:
            fail(f"{xlsx_path.name} 的 {get_column_letter(col)}{row} 现值是 {current!r},不是预期的 {old!r} —— 未写任何文件。")
        if new is None:
            fail(f"内部错误:改写的新值为空({get_column_letter(col)}{row})")

    if loaded_digest and file_digest(xlsx_path) != loaded_digest:
        fail(f"{xlsx_path.name} 在本脚本载入之后被别的进程改过(并发写),保存会抹掉它的改动 —— 未写任何文件,重跑本脚本即可。")
    for row, col, value in patches:
        ws.cell(row=row, column=col, value=value)
    for row, col, _old, new in rewrites:
        ws.cell(row=row, column=col, value=new)
    wb.save(xlsx_path)

    touched = [(r, c) for r, c, _v in patches] + [(r, c) for r, c, _o, _n in rewrites]
    width = max([len(r) for r in before] + [c for _r, c in touched])
    height = max([len(before)] + [r for r, _c in touched])
    expected = [list(r) + [None] * (width - len(r)) for r in before]
    expected += [[None] * width for _ in range(height - len(expected))]
    for row, col, value in patches:
        expected[row - 1][col - 1] = value
    for row, col, _old, new in rewrites:
        expected[row - 1][col - 1] = new

    check_ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path, sheet)
    diff = first_difference(sheet_matrix(check_ws), expected)
    if diff:
        fail(f"保存后读回的 {xlsx_path.name} 与预期不符({diff})—— 立刻 git checkout -- {rel(xlsx_path)} 还原")


# ════════════════════════════════════════════════════════════════════════════
# ① tip-segment
# ════════════════════════════════════════════════════════════════════════════

TIP_XLSX = DATA_DIR / "tip" / "Tip.xlsx"
TIP_STATE = EXPORTER_DIR / "state" / "mapping" / "tip_enum_ids" / "tip_enum_ids.json"
TIP_DATA_BEGIN_ROW = 18            # 与导表器 enum_gen._TIP_DATA_BEGIN_ROW 一致
TIP_PREFIX = "//"


def is_group_header(name: str) -> bool:
    """组头 vs 说明行,判据与导表器一致(enum_gen._parse_tip_sheet):"//" 后**紧贴**非空白字符 = 组头,
    "// 这是一句说明" = 说明行。不能用"有没有 base="来分,那会与导表器的口径分叉。"""
    return name.startswith(TIP_PREFIX) and len(name) > 2 and not name[2].isspace()


def parse_group_header(name: str, row_idx: int) -> tuple[str, int, int]:
    parts = name[2:].split()
    kv = dict(tok.split("=", 1) for tok in parts[1:] if "=" in tok)
    try:
        return parts[0], int(kv["base"]), int(kv["width"])
    except (KeyError, ValueError):
        fail(f"{TIP_XLSX.name} 第 {row_idx} 行 {name!r} 像组头却没有合法的 base= / width= —— 导表器也会拒绝,先修表")


def tip_parse(matrix: list[list], fault_col: int):
    """返回 (组列表, 说明行列表)。组 = dict(name, base, width, row, codes=[(码名, 文案, fault 原值)])。"""
    groups: list[dict] = []
    notes: list[tuple[int, str]] = []
    for row_idx in range(TIP_DATA_BEGIN_ROW, len(matrix) + 1):
        row = matrix[row_idx - 1]
        a = row[0]
        if not isinstance(a, str) or not a.strip():
            continue
        name = a.strip()
        if name.startswith(TIP_PREFIX):
            if not is_group_header(name):
                notes.append((row_idx, name))
                continue
            group, base, width = parse_group_header(name, row_idx)
            groups.append(dict(name=group, base=base, width=width, row=row_idx, codes=[]))
            continue
        if not groups:
            fail(f"{TIP_XLSX.name} 第 {row_idx} 行的码 {name!r} 出现在第一个组头之前")
        fault = row[fault_col - 1] if fault_col <= len(row) else None
        groups[-1]["codes"].append((name, row[1] if len(row) > 1 else None, fault))
    return groups, notes


def reservation_rewrites(matrix: list[list], notes: list[tuple[int, str]], group_exists: bool) -> list[Rewrite]:
    """预留说明那一行要不要改写。

    mail 组还没建(group_exists=False)时最严:提到 17000 的说明行只允许三种现状 —— ① 逐字等于旧文案(要改写);
    ② 逐字等于新文案(已改过);③ 根本没有(有人整理过那几行)。其它任何提到 17000 的说明都中止:
    那可能是别的域把这一段预留走了。旧、新文案同时出现也中止(改写后会有两行一样的新文案)。

    mail 组已经建好(group_exists=True)时 17000 段已是本组的,别的说明行再提到它不构成威胁(比如 chat 建组时
    又改写了这一行,新文案里照样带着 17000):只看还有没有逐字等于旧文案的行,有就改写,没有就什么都不做。
    """
    if group_exists:
        olds = [row for row, _text in notes if matrix[row - 1][0] == TIP_RESERVATION_OLD]
        if len(olds) > 1:
            fail(f"{TIP_XLSX.name} 里有 {len(olds)} 行相同的旧预留说明(第 {olds} 行)—— 先修表")
        return [(row, 1, TIP_RESERVATION_OLD, TIP_RESERVATION_NEW) for row in olds]

    mentioning = [(row, text) for row, text in notes if str(TIP_MAIL_BASE) in text]
    if any(t == TIP_RESERVATION_OLD for _r, t in mentioning) and any(t == TIP_RESERVATION_NEW for _r, t in mentioning):
        fail(f"{TIP_XLSX.name} 里旧、新两版预留说明同时存在(第 {[r for r, _t in mentioning]} 行)—— 先修表,改写后会出现两行相同的说明")
    rewrites: list[Rewrite] = []
    for row, text in mentioning:
        if text == TIP_RESERVATION_NEW:
            continue
        if text != TIP_RESERVATION_OLD:
            fail(f"{TIP_XLSX.name} 第 {row} 行的说明 {text!r} 提到了 {TIP_MAIL_BASE},但不是本脚本认识的那一行预留说明 —— "
                 f"先确认 {TIP_MAIL_BASE} 段是不是被别的域预留了。")
        raw = matrix[row - 1][0]
        if raw != TIP_RESERVATION_OLD:
            # 去掉首尾空白后才相等:照原样改写会把别人留的空白抹掉,这里宁可不动。
            fail(f"{TIP_XLSX.name} 第 {row} 行的预留说明首尾带空白({raw!r}),本脚本不改写 —— 手工确认后再跑。")
        rewrites.append((row, 1, TIP_RESERVATION_OLD, TIP_RESERVATION_NEW))
    if len(rewrites) > 1:
        fail(f"{TIP_XLSX.name} 里有 {len(rewrites)} 行相同的预留说明(第 {[r for r, *_ in rewrites]} 行)—— 先修表")
    return rewrites


def cmd_tip_segment(dry_run: bool) -> None:
    loaded_digest = file_digest(TIP_XLSX)
    wb = openpyxl.load_workbook(TIP_XLSX)
    ws = only_sheet(wb, TIP_XLSX)
    before = sheet_matrix(ws)
    header = [str(v).strip().lower() if v is not None else "" for v in before[HEADER_ROW - 1]]
    if header[:2] != ["name", "commen"] or header.count("fault") != 1:
        fail(f"{TIP_XLSX.name} 第 1 行应为 name / commen / … / fault(fault 恰好一列),实际 {before[HEADER_ROW - 1]}")
    fault_col = header.index("fault") + 1

    groups, notes = tip_parse(before, fault_col)
    header_text = f"//{TIP_MAIL_GROUP} base={TIP_MAIL_BASE} width={TIP_MAIL_WIDTH}"
    wanted = [(name, text, None) for name, text in TIP_MAIL_CODES]      # fault 列全空
    new_names = [name for name, _t in TIP_MAIL_CODES]
    mine = next((g for g in groups if g["name"] == TIP_MAIL_GROUP), None)
    rewrites = reservation_rewrites(before, notes, mine is not None)

    patches: list[Patch] = []
    if mine is not None:
        # 已是目标状态 = 段声明一致,且这 8 个码是本组开头的连续一段。不要求恰好在组尾:M2 还会往组尾追加。
        if (mine["base"], mine["width"]) != (TIP_MAIL_BASE, TIP_MAIL_WIDTH) or mine["codes"][:len(wanted)] != wanted:
            fail(f"{TIP_XLSX.name} 第 {mine['row']} 行已有 //{TIP_MAIL_GROUP} 组,但段声明或开头 {len(wanted)} 个码"
                 f"(码名 / 文案 / fault 留空)与本脚本不同 —— 先查清是谁加的。")
        if not rewrites:
            print(f"[ALREADY CURRENT] {header_text} (row {mine['row']}) already contains these {len(wanted)} codes "
                  f"and the reservation note is up to date; file unchanged.")
            return
        print(f"{TIP_XLSX.name}: {header_text} (row {mine['row']}) already present; only the reservation note needs updating.")
    else:
        # 以下几项都是「先核对再动手」:码名全表唯一、17000 段没被占、发号 state 里也没有这一段的墓碑。
        owner = {code: g["name"] for g in groups for code, _t, _f in g["codes"]}
        clash = [f"{n}(在 {owner[n]})" for n in new_names if n in owner]
        if clash:
            fail(f"这些码名已在表里:{clash} —— 码名全局唯一(生成的 tip proto 没有 package),先查清。")
        lo, hi = TIP_MAIL_BASE, TIP_MAIL_BASE + TIP_MAIL_WIDTH
        overlap = [f"{g['name']} [{g['base']},{g['base'] + g['width']})" for g in groups
                   if g["base"] < hi and lo < g["base"] + g["width"]]
        if overlap:
            fail(f"段 [{lo},{hi}) 已被占用:{overlap} —— {TIP_MAIL_BASE} 是设计文档 §5.6 定下的段,"
                 f"被占说明有人抢了号,先查清,不要自行换段。")
        if TIP_STATE.exists():
            # 发号 state 里的整组是墓碑:表里删掉的组仍占着段。只读核对,不写 state。
            # 唯一放行的墓碑是**本组自己的**:导表之后 Tip.xlsx 被别的会话整表覆盖、mail 组丢了(本脚本头注释说的那类事故),
            # state 里仍记着 mail_error 占 17000 段。段声明相同、已发号的码名都在本脚本清单里时,补回这一组就是恢复原状,
            # 导表器会按 state 给它们发回原来的号。
            state_groups = json.loads(TIP_STATE.read_text(encoding="utf-8")).get("groups", {})
            tomb = []
            for n, g in state_groups.items():
                if not (int(g["base"]) < hi and lo < int(g["base"]) + int(g["width"])):
                    continue
                ours = (n == TIP_MAIL_GROUP and (int(g["base"]), int(g["width"])) == (TIP_MAIL_BASE, TIP_MAIL_WIDTH)
                        and set(g.get("ids", {})) <= set(new_names))
                if ours:
                    print(f"{TIP_XLSX.name}: issued-id state already holds {TIP_MAIL_GROUP} [{lo},{hi}) with codes from this script;"
                          f" the group is missing from the sheet (overwritten?) and will be restored.")
                    continue
                tomb.append(f"{n} [{g['base']},{int(g['base']) + int(g['width'])})")
            if tomb:
                fail(f"发号 state({rel(TIP_STATE)})里段 [{lo},{hi}) 已属于 {tomb} —— 不要手改 state,先查清。")

        last_row = max((r for r, row in enumerate(before, start=1) if any(v is not None for v in row)), default=0)
        header_row = last_row + 2          # 与既有组一样,组头前留一个空行
        patches.append((header_row, 1, header_text))
        print(f"{TIP_XLSX.name}: existing {len(groups)} groups, {len(owner)} codes; last row={last_row};"
              f" range [{lo},{hi}) available; {len(new_names)} code names globally unique. Will append:")
        print(f"  Row {header_row}: {header_text}")
        for offset, (name, text, _fault) in enumerate(wanted, start=1):
            row_idx = header_row + offset
            patches += [(row_idx, 1, name), (row_idx, 2, text)]
            print(f"  Row {row_idx}: {name} | {text}")

    for row, _col, old, new in rewrites:
        print(f"  Row {row} (reservation note): {old!r} -> {new!r}")
    if dry_run:
        print("[dry-run] No files written.")
        return
    apply_and_verify(TIP_XLSX, None, wb, ws, before, patches, rewrites, loaded_digest)

    check_ws = only_sheet(openpyxl.load_workbook(TIP_XLSX), TIP_XLSX)
    groups_after, _notes = tip_parse(sheet_matrix(check_ws), fault_col)
    others_before = [(g["name"], g["base"], g["width"], g["codes"]) for g in groups if g["name"] != TIP_MAIL_GROUP]
    others_after = [(g["name"], g["base"], g["width"], g["codes"]) for g in groups_after if g["name"] != TIP_MAIL_GROUP]
    mine_after = [g for g in groups_after if g["name"] == TIP_MAIL_GROUP]
    if others_after != others_before or len(mine_after) != 1 or mine_after[0]["codes"][:len(wanted)] != wanted:
        fail(f"保存后各组的码与预期不符 —— 立刻 git checkout -- {rel(TIP_XLSX)} 还原")
    lo = TIP_MAIL_BASE
    print(f"[DONE] Groups {len(groups)} -> {len(groups_after)}; all other groups unchanged (names / ranges / codes)."
          f" After export, confirm these {len(wanted)} codes have IDs {lo}..{lo + len(wanted) - 1}; tip_enum_ids.json contains additions only;"
          f" faults.go has no {lo // 1000}xxx code;"
          f" register proto/tip/{TIP_MAIL_GROUP}_tip.pb.{{h,cc}} in cpp/generated/table CMakeLists.txt, table.vcxproj and table.vcxproj.filters.")


# ════════════════════════════════════════════════════════════════════════════
# ② message-limiter
# ════════════════════════════════════════════════════════════════════════════

MESSAGE_ID_TXT = REPO_ROOT / "proto" / "message_id.txt"
LIMITER_SHEET = "MessageLimiter"
LIMITER_XLSX = DATA_DIR / "MessageLimiter.xlsx"
LIMITER_COLUMNS = ("id", "max_requests", "time_window", "tip_message")


def resolve_message_ids() -> dict[str, int]:
    """方法名 -> 消息号。任何一个方法查不到就中止 —— 此时还没有打开过 xlsx,什么都不会写。"""
    if not MESSAGE_ID_TXT.is_file():
        fail(f"找不到 {rel(MESSAGE_ID_TXT)} —— 消息号表由全量 proto-gen 生成,先跑 proto-gen。未写任何文件。")
    ids: dict[str, int] = {}
    owner: dict[int, str] = {}
    for line_no, line in enumerate(MESSAGE_ID_TXT.read_text(encoding="utf-8").splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        number, sep, key = line.partition("=")
        if not sep or not number.isdigit():
            fail(f"{MESSAGE_ID_TXT.name}:{line_no} 不是 `<号>=<服务名+方法名>` 形状:{line!r}")
        if key in ids:
            fail(f"{MESSAGE_ID_TXT.name}:{line_no} 的 {key!r} 出现了两次 —— 消息号表坏了,先修。未写任何文件。")
        if int(number) in owner:
            fail(f"{MESSAGE_ID_TXT.name}:{line_no} 的号 {number} 同时分给了 {owner[int(number)]} 与 {key}"
                 f" —— 消息号表坏了,先修。未写任何文件。")
        ids[key] = int(number)
        owner[int(number)] = key
    needed = [LIMITER_SERVICE + m for m in list(LIMITER_METHODS) + [LIMITER_S2C_METHOD]]
    missing = [key for key in needed if key not in ids]
    if missing:
        fail(f"{MESSAGE_ID_TXT.name} 里找不到 {missing} —— 含 mail 的全量 proto-gen 还没跑(或没跑成功),消息号未定,"
             f"现在不能填档位。未写任何文件。")
    return {m: ids[LIMITER_SERVICE + m] for m in LIMITER_METHODS}


def header_cells(ws) -> dict[str, int]:
    """第 1 行非空表头 -> {列名: 1 起的列号}。重名直接中止(导表器也会拒绝)。"""
    out: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=HEADER_ROW, column=c).value
        name = str(value).strip() if value is not None else ""
        if not name:
            continue
        if name in out:
            fail(f"{ws.title} 第 1 行有重复表头 {name!r}")
        out[name] = c
    return out


def bind_limiter_schema(after_write: bool) -> None:
    """用导表器自己的绑定器核对 schema 与 xlsx:本脚本说"好了"与导表器说"好了"是同一个判据。

    写之前先绑一次(after_write=False):绑不上就什么都不写 —— 否则表已改、核对却失败,重跑时又走
    "已是目标状态"分支、不再核对,失败被掩盖。写之后再绑一次,失败时提示还原。
    """
    state = (f"{rel(LIMITER_XLSX)} 已写入 —— 先 git checkout -- {rel(LIMITER_XLSX)} 还原再查"
             if after_write else "未写任何文件")
    if str(EXPORTER_DIR) not in sys.path:
        sys.path.insert(0, str(EXPORTER_DIR))
    try:
        from core.schema_proto import SchemaProtoError, build_table_schema
    except ImportError as exc:  # pragma: no cover - 环境问题
        fail(f"无法载入导表器的 schema 解析器({exc}):请先执行  pip install pyyaml openpyxl;{state}")
    try:
        build_table_schema(SCHEMA_DIR / f"{LIMITER_SHEET.lower()}_table.proto", LIMITER_XLSX)
    except SchemaProtoError as exc:
        fail(f"权威 schema 与 {rel(LIMITER_XLSX)} 对不上(导表也会因此失败):{exc};{state}")


def cmd_message_limiter(dry_run: bool) -> None:
    message_ids = resolve_message_ids()
    bind_limiter_schema(after_write=False)

    loaded_digest = file_digest(LIMITER_XLSX)
    wb = openpyxl.load_workbook(LIMITER_XLSX)
    ws = only_sheet(wb, LIMITER_XLSX, LIMITER_SHEET)
    before = sheet_matrix(ws)
    header = header_cells(ws)
    for name in LIMITER_COLUMNS:
        if name not in header:
            fail(f"{LIMITER_XLSX.name} 第 1 行找不到列 {name!r},实际表头 {list(header)}")

    existing: dict[int, tuple] = {}
    last_data_row = DATA_BEGIN_ROW - 1
    for row_idx in range(DATA_BEGIN_ROW, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=header["id"]).value
        if value is None or value == "":
            continue
        if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
            fail(f"{LIMITER_XLSX.name} 第 {row_idx} 行的 id 不是整数:{value!r}")
        if int(value) in existing:
            fail(f"{LIMITER_XLSX.name} 里消息号 {int(value)} 出现了两次(第 {row_idx} 行)—— 先修表")
        existing[int(value)] = tuple(ws.cell(row=row_idx, column=header[name]).value for name in LIMITER_COLUMNS[1:])
        last_data_row = row_idx

    to_append = []
    for method, max_requests in LIMITER_METHODS.items():
        message_id = message_ids[method]
        wanted = (max_requests, LIMITER_TIME_WINDOW, LIMITER_TIP_MESSAGE)
        if message_id not in existing:
            to_append.append((message_id, method, wanted))
        elif existing[message_id] != wanted:
            # 最可能的原因:这个号以前属于别的方法(被释放的旧号易主),表里留着旧方法的档位。
            fail(f"消息号 {message_id}({LIMITER_SERVICE}{method})在表里已有一行 {existing[message_id]},"
                 f"与要写的 {wanted} 不同 —— 先确认那一行是谁的,本脚本不覆盖。未写任何文件。")

    if not to_append:
        print(f"[ALREADY CURRENT] All {len(message_ids)} mail C2S message IDs exist in the table; file unchanged.")
        return

    patches: list[Patch] = []
    print(f"Existing data: {len(existing)} rows; last row={last_data_row}; will append {len(to_append)} rows:")
    for offset, (message_id, method, wanted) in enumerate(to_append, start=1):
        row_idx = last_data_row + offset
        patches.append((row_idx, header["id"], message_id))
        patches += [(row_idx, header[name], value) for name, value in zip(LIMITER_COLUMNS[1:], wanted)]
        print(f"  Row {row_idx}: id={message_id:<4} {LIMITER_SERVICE}{method:<22} "
              f"max_requests={wanted[0]} time_window={wanted[1]} tip_message={wanted[2]}")
    if dry_run:
        print("[dry-run] No files written.")
        return
    apply_and_verify(LIMITER_XLSX, LIMITER_SHEET, wb, ws, before, patches, (), loaded_digest)
    bind_limiter_schema(after_write=True)
    print(f"[DONE] Data rows {len(existing)} -> {len(existing) + len(to_append)}; existing cells unchanged. Rerun the table exporter;"
          f" acceptance: the IDs above appear in generated/tables/messagelimiter.json. Reload gate MessageLimiter during deployment.")


COMMANDS = {
    "tip-segment": cmd_tip_segment,
    "message-limiter": cmd_message_limiter,
}


def self_check() -> None:
    """常量表自检:写错常量比写错表更难发现,先在这里拦。"""
    names = [name for name, _t in TIP_MAIL_CODES]
    if len(set(names)) != len(names):
        fail("内部错误:TIP_MAIL_CODES 里有重复码名")
    for name, text in TIP_MAIL_CODES:
        if not name.startswith("Mail") or not name.isidentifier() or not name.isascii():
            fail(f"内部错误:码名 {name!r} 应当是 Mail 开头的 ASCII 标识符(码名全局唯一,靠域前缀防撞)")
        if not text or text != text.strip():
            fail(f"内部错误:码 {name} 的文案为空或首尾带空白")
    if len(TIP_MAIL_CODES) > TIP_MAIL_WIDTH:
        fail("内部错误:码数超过段宽")
    if is_group_header(TIP_RESERVATION_NEW) or is_group_header(TIP_RESERVATION_OLD):
        fail("内部错误:预留说明会被导表器当成组头(// 后必须跟空白)")
    if str(TIP_MAIL_BASE) not in TIP_RESERVATION_OLD:
        fail("内部错误:旧预留说明里应当提到本段的起始号")
    for method, tier in LIMITER_METHODS.items():
        if not isinstance(tier, int) or tier <= 0:
            fail(f"内部错误:{method} 的档位必须是正整数")
    if LIMITER_S2C_METHOD in LIMITER_METHODS:
        fail("内部错误:S2C 方法不进限流表")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=tuple(COMMANDS))
    parser.add_argument("--dry-run", action="store_true", help="只打印将要做什么,不写文件")
    args = parser.parse_args()
    self_check()
    COMMANDS[args.command](args.dry_run)


if __name__ == "__main__":
    # Windows 默认控制台编码(cp936 / cp1252)打不出部分字符时会在 print 上抛 UnicodeEncodeError,
    # 而那时文件可能已经写完 —— 统一成 utf-8,免得"改成功了却以非 0 退出"误导人。
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    main()
