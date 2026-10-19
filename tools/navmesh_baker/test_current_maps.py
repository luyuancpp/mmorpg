"""新地图跨端合同的失配拒绝与旧场景保留回归。"""
import base64
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from sync_current_maps import digest, expected_rows, load_contract


class CurrentMapContractTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="current-map-contract-")
        self.folder = Path(self.temporary.name)
        mask = bytearray(2813)
        # (199,0,133) => col74,row83,必须从上到下/MSB-first读取。
        index = 83 * 150 + 74
        mask[index // 8] |= 0x80 >> (index % 8)
        self.mask = bytes(mask)
        (self.folder / "walkmask.txt").write_bytes(base64.b64encode(mask))
        (self.folder / "layout.png").write_bytes(b"test-art-bytes")
        (self.folder / "navigation.json").write_bytes(b"{}")
        self.source = {
            "schemaVersion": 1, "sceneId": 1, "key": "main",
            "walkmaskResolution": 150, "cellSize": 2,
            "worldRect": {"x": 50, "y": 0, "width": 300, "height": 300},
            "maskEncoding": "base64-msb-first-row-major-top-to-bottom",
            "spawnWorld": {"x": 199, "y": 0, "z": 133},
            "walkableCellCount": 1,
            "walkmaskSha256": digest(base64.b64encode(mask)),
            "walkmaskPackedSha256": digest(mask),
            "layoutSha256": digest(b"test-art-bytes"),
            "navigationSha256": digest(b"{}"),
        }

    def tearDown(self):
        self.temporary.cleanup()

    def write_source(self):
        (self.folder / "source.json").write_text(json.dumps(self.source), encoding="utf-8")

    def test_real_coordinate_and_bit_order_contract(self):
        self.write_source()
        self.assertEqual(self.source, load_contract(self.folder, 1, "main"))

    def test_modified_mask_fails_before_baking(self):
        self.write_source()
        (self.folder / "walkmask.txt").write_bytes(base64.b64encode(bytes(2813)))
        with self.assertRaisesRegex(ValueError, "SHA256"):
            load_contract(self.folder, 1, "main")

    def test_blocked_or_nonfinite_spawn_is_rejected(self):
        for spawn in ({"x": 197, "y": 0, "z": 133}, {"x": float("nan"), "y": 0, "z": 133}):
            self.source["spawnWorld"] = spawn
            self.write_source()
            with self.assertRaises(ValueError):
                load_contract(self.folder, 1, "main")

    def test_unknown_geometry_or_scene_identity_is_rejected(self):
        for field, value in (("cellSize", 1), ("sceneId", 4), ("walkmaskResolution", 300)):
            with self.subTest(field=field):
                old = self.source[field]
                self.source[field] = value
                self.write_source()
                with self.assertRaises(ValueError):
                    load_contract(self.folder, 1, "main")
                self.source[field] = old

    def test_shared_nav_rows_follow_spawn_but_other_maps_stay_unchanged(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["id", "nav_bin_file", "spawn_x", "spawn_y", "spawn_z"])
        for _ in range(4):
            sheet.append([None])
        for scene_id, filename in ((1, "main"), (2, "penglai"), (3, "donghai"),
                                   (4, "lanxian"), (5, "main"), (17, "dungeon")):
            sheet.append([scene_id, f"data/scene_nav_bin/{filename}_scene.bin", 180, 200, 0])
        contracts = {key: {"spawnWorld": {"x": x, "y": 0, "z": z}}
                     for key, x, z in (("main", 199, 133), ("island", 195, 173), ("village", 201, 147))}
        rows, changes = expected_rows(book, contracts)
        by_id = {row["id"]: row for row in rows}
        self.assertEqual((133, 199, 0), tuple(by_id[1][f"spawn_{a}"] for a in "xyz"))
        self.assertEqual((173, 195, 0), tuple(by_id[2][f"spawn_{a}"] for a in "xyz"))
        self.assertEqual((147, 201, 0), tuple(by_id[3][f"spawn_{a}"] for a in "xyz"))
        self.assertEqual(by_id[1]["spawn_x"], by_id[5]["spawn_x"])
        for scene_id in (4, 17):
            self.assertEqual((180, 200, 0), tuple(by_id[scene_id][f"spawn_{a}"] for a in "xyz"))
        self.assertTrue(changes)
        self.assertEqual(180, sheet.cell(6, 3).value, "预检本身不得修改源表")
        sheet.cell(6, 2).value = "data/scene_nav_bin/lanxian_scene.bin"
        with self.assertRaisesRegex(ValueError, "实际导航路径"):
            expected_rows(book, contracts)
        book.close()


if __name__ == "__main__":
    unittest.main()
