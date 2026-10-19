#!/usr/bin/env python3
"""从客户端当前地图合同同步正式导航与 BaseScene 源表;不修改生成表。

同步后走现有导表器生成 BaseScene JSON/PB 与 manifest,再用 --verify 验证。
默认只替换三张正式地图与复用同一导航文件的出生点,不改其它地图、账号或数据库。
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
MAPS = ((1, "main", "main_scene.bin"), (2, "island", "penglai_scene.bin"),
        (3, "village", "donghai_scene.bin"))
MANIFEST = "current_maps_20261009.json"


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_contract(folder: Path, scene_id: int, key: str) -> dict:
    source = json.loads((folder / "source.json").read_text(encoding="utf-8"))
    expected = {"schemaVersion": 1, "sceneId": scene_id, "key": key,
                "walkmaskResolution": 150, "cellSize": 2,
                "worldRect": {"x": 50, "y": 0, "width": 300, "height": 300},
                "maskEncoding": "base64-msb-first-row-major-top-to-bottom"}
    for field, value in expected.items():
        if source.get(field) != value:
            raise ValueError(f"{key}: 不支持的地图合同 {field}={source.get(field)!r}")
    for filename, sha_field in (("walkmask.txt", "walkmaskSha256"),
                                ("layout.png", "layoutSha256"),
                                ("navigation.json", "navigationSha256")):
        if digest((folder / filename).read_bytes()) != source.get(sha_field):
            raise ValueError(f"{key}: {filename} 与 source.json 的 SHA256 不一致")
    packed = base64.b64decode((folder / "walkmask.txt").read_bytes().strip(), validate=True)
    if len(packed) != 2813 or packed[-1] & 15:
        raise ValueError(f"{key}: 位图必须包含 22500 格,尾部填充位必须为零")
    if digest(packed) != source.get("walkmaskPackedSha256"):
        raise ValueError(f"{key}: 解码后位图 SHA256 不一致")
    if sum(byte.bit_count() for byte in packed) != source.get("walkableCellCount"):
        raise ValueError(f"{key}: 可走格数量不一致")
    point = source.get("spawnWorld", {})
    spawn = [point.get(axis) for axis in ("x", "y", "z")]
    if any(type(value) not in (int, float) or not math.isfinite(value) for value in spawn):
        raise ValueError(f"{key}: 出生点必须是有限坐标")
    x, y, z = spawn
    if not (50 <= x < 350 and 0 < z <= 300 and y == 0):
        raise ValueError(f"{key}: 出生点超出位图地面")
    col, row = int((x - 50) // 2), int((300 - z) // 2)
    index = row * 150 + col
    if not packed[index // 8] & (0x80 >> (index % 8)):
        raise ValueError(f"{key}: 出生点在不可走格中")
    return source


def expected_rows(workbook, contracts: dict) -> tuple[list, list]:
    sheet = workbook.active
    columns = {cell.value: cell.column for cell in sheet[1]}
    required = ("id", "nav_bin_file", "spawn_x", "spawn_y", "spawn_z")
    if set(columns) != set(required):
        raise ValueError("BaseScene 表头发生变化,请先检查 schema")
    paths = {f"data/scene_nav_bin/{filename}": contracts[key]
             for _, key, filename in MAPS}
    changes, rows = [], []
    seen = {}
    for index in range(6, sheet.max_row + 1):
        row = {field: sheet.cell(index, columns[field]).value for field in required}
        if row["id"] is None:
            continue
        if row["id"] in seen:
            raise ValueError(f"BaseScene 重复 ID {row['id']}")
        seen[row["id"]] = row["nav_bin_file"]
        source = paths.get(row["nav_bin_file"])
        if source:
            point = source["spawnWorld"]
            # Unity (x,y,z) => 服务端 Z-up (z,x,y)。同文件所有行共享出生合同。
            for field, value in zip(("spawn_x", "spawn_y", "spawn_z"),
                                    (point["z"], point["x"], point["y"])):
                if row[field] != value:
                    changes.append((index, columns[field], value))
                row[field] = value
        rows.append(row)
    for scene_id, _, filename in MAPS:
        if seen.get(scene_id) != f"data/scene_nav_bin/{filename}":
            raise ValueError(f"场景 {scene_id} 的实际导航路径改变,拒绝覆盖错误文件")
    return rows, changes


def verify_generated(repo: Path, expected: list) -> None:
    table_dir = repo / "generated/tables"
    actual = json.loads((table_dir / "basescene.json").read_text(encoding="utf-8"))["data"]
    if actual != expected:
        raise ValueError("BaseScene 生成表与新地图不一致,请先运行正式导表器")
    manifest = json.loads((table_dir / "manifest.json").read_text(encoding="utf-8"))
    table = next(item for item in manifest["tables"] if item["name"] == "BaseScene")
    checks = [(repo / "data" / table["source"]["file"], table["source"])]
    checks.extend((table_dir / item["file"], item) for item in table["artifacts"])
    for path, item in checks:
        data = path.read_bytes()
        if digest(data) != item["sha256"] or len(data) != item["size"]:
            raise ValueError(f"正式导表清单与文件不一致: {path}")


def sync(repo: Path, source_root: Path, baker: Path, verify: bool = False) -> dict:
    contracts = {key: load_contract(source_root / key, scene_id, key)
                 for scene_id, key, _ in MAPS}
    snapshots = {key: {filename: (source_root / key / filename).read_bytes()
                       for filename in ("source.json", "walkmask.txt")}
                 for _, key, _ in MAPS}
    for key, files in snapshots.items():
        if (json.loads(files["source.json"]) != contracts[key] or
                digest(files["walkmask.txt"]) != contracts[key]["walkmaskSha256"]):
            raise ValueError(f"{key}: 预检期间客户端合同发生变化,请在资源稳定后重新同步")
    xlsx = repo / "data/BaseScene.xlsx"
    workbook = openpyxl.load_workbook(xlsx)
    rows, changes = expected_rows(workbook, contracts)
    if verify and changes:
        raise ValueError("BaseScene.xlsx 仍含旧出生点,请先同步地图")
    records = []
    # 三图全部烘焙/探针成功后才替换正式文件,第三张失败不会提前写前两张。
    with tempfile.TemporaryDirectory(prefix="current-map-nav-") as temporary:
        temporary = Path(temporary)
        for scene_id, key, filename in MAPS:
            source = contracts[key]
            point = source["spawnWorld"]
            target = temporary / filename
            mask_file = temporary / f"{key}-walkmask.txt"
            mask_file.write_bytes(snapshots[key]["walkmask.txt"])
            command = [str(baker), "--mask-base64", str(mask_file),
                       "--probe", ",".join(str(point[axis]) for axis in ("x", "y", "z")),
                       "--out", str(target)]
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", timeout=120)
            print(result.stdout, end="")
            if result.returncode != 0:
                raise RuntimeError(f"{key} 烘焙失败: {result.stderr}")
            if "mask cells checked=22500 mismatches=0" not in result.stdout:
                raise ValueError("烘焙器版本不支持逐格校验,请先重新构建")
            nav = target.read_bytes()
            records.append({"sceneId": scene_id, "key": key, "name": source["name"],
                            "sourceResource": f"World/CurrentMaps20261009/{key}",
                            "layoutSha256": source["layoutSha256"],
                            "walkmaskSha256": source["walkmaskSha256"],
                            "walkmaskPackedSha256": source["walkmaskPackedSha256"],
                            "walkableCellCount": source["walkableCellCount"],
                            "spawnUnity": point,
                            "spawnServer": {"x": point["z"], "y": point["x"], "z": point["y"]},
                            "navPath": f"data/scene_nav_bin/{filename}",
                            "navSha256": digest(nav), "navBytes": len(nav),
                            "checkedCellCount": 22500, "cellMismatchCount": 0})
        manifest = {"schemaVersion": 1, "maps": records}
        encoded = (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        nav_dir = repo / "data/scene_nav_bin"
        if verify:
            for _, _, filename in MAPS:
                if (nav_dir / filename).read_bytes() != (temporary / filename).read_bytes():
                    raise ValueError(f"{filename} 与当前客户端位图重新烘焙的结果不一致")
            if (nav_dir / MANIFEST).read_bytes() != encoded:
                raise ValueError("导航来源清单与当前地图不一致")
            for _, key, _ in MAPS:
                for filename in ("source.json", "walkmask.txt"):
                    if (nav_dir / "current_maps" / key / filename).read_bytes() != snapshots[key][filename]:
                        raise ValueError(f"{key}/{filename} 的跨语言合同副本不一致")
            verify_generated(repo, rows)
        else:
            for _, _, filename in MAPS:
                shutil.copyfile(temporary / filename, nav_dir / filename)
            # Java 导航读取同一位图合同;这是客户端事实源的逐字节产物,不可单独编辑。
            for _, key, _ in MAPS:
                folder = nav_dir / "current_maps" / key
                folder.mkdir(parents=True, exist_ok=True)
                for filename in ("source.json", "walkmask.txt"):
                    (folder / filename).write_bytes(snapshots[key][filename])
            for row, column, value in changes:
                workbook.active.cell(row, column).value = value
            if changes:
                workbook.save(xlsx)
            (nav_dir / MANIFEST).write_bytes(encoded)
    workbook.close()
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--client-root", type=Path, default=ROOT.parent / "mmorpg-client")
    parser.add_argument("--baker", type=Path, default=ROOT / "tools/navmesh_baker/build/Release/navmesh_baker.exe")
    parser.add_argument("--verify", action="store_true", help="只验证资源、重新烘焙一致性和正式导表清单")
    args = parser.parse_args()
    source_root = args.client_root.resolve() / "Assets/Resources/World/CurrentMaps20261009"
    result = sync(ROOT, source_root, args.baker.resolve(), args.verify)
    print(f"{'验证' if args.verify else '同步'}通过: {len(result['maps'])} 张地图,67500 格导航一致")
    if not args.verify:
        print("接下来运行正式导表器,再用本工具 --verify 核对 BaseScene JSON/PB 与 manifest。")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
